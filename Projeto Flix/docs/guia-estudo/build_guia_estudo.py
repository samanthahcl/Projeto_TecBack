from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent / "output"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_DOCX = OUT_DIR / "IESPFLIX_Guia_Estudo_Atualizado.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK_BLUE = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
MUTED = "666666"
WHITE = "FFFFFF"
CODE_BG = "F5F7FA"
CODE_BORDER = "D9E1EA"
GREEN = "1F6D42"
GOLD = "7A5A00"
RED = "9B1C1C"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:{}".format(key)), str(value))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths_dxa):
        col.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def setup_numbering(doc):
    numbering = doc.part.numbering_part.element

    def create_abstract(num_id, fmt, text, left=540, hanging=270, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(num_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{num_id:08X}")
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font)
            fonts.set(qn("w:hAnsi"), font)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def create_num(num_id, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abstract_id))
        num.append(abstract_num_id)
        numbering.append(num)

    create_abstract(90, "bullet", "•", font="Symbol")
    create_abstract(91, "decimal", "%1.")
    create_num(90, 90)
    create_num(91, 91)
    return 90, 91


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def add_bullet(doc, text, bullet_num, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num)
    set_paragraph_spacing(p, after=4, line=1.25)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_number(doc, text, decimal_num):
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    set_paragraph_spacing(p, after=4, line=1.25)
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p


def add_para(doc, text="", *, bold=False, italic=False, color=None, align=None, before=0, after=6, line=1.25):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=line)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, bold=bold, italic=italic, color=color)
    return p


def add_rich_para(doc, parts, *, before=0, after=6, line=1.25):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=line)
    for text, options in parts:
        run = p.add_run(text)
        set_font(
            run,
            name=options.get("font", "Calibri"),
            size=options.get("size", 11),
            color=options.get("color"),
            bold=options.get("bold"),
            italic=options.get("italic"),
        )
    return p


def add_code(doc, code):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=0, line=1.0)
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.right_indent = Inches(0.08)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), CODE_BG)
        p_pr.append(shd)
        run = p.add_run(line if line else " ")
        set_font(run, name="Consolas", size=8.2, color=INK_BLUE)
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=4, line=1.0)


def add_callout(doc, title, text, fill=CALLOUT, title_color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    border = {"val": "single", "sz": "4", "space": "0", "color": "D9E1EA"}
    set_cell_border(cell, top=border, left=border, bottom=border, right=border)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2, line=1.15)
    r = p.add_run(title)
    set_font(r, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=0, line=1.15)
    r2 = p2.add_run(text)
    set_font(r2, size=10.5)
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=4, line=1.0)


def add_table(doc, headers, rows, widths, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_paragraph_spacing(p, after=0, line=1.0)
        r = p.add_run(str(value))
        set_font(r, size=font_size, bold=True, color=INK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.05)
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=4, line=1.0)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    style_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in style_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15


def configure_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(hp, after=0, line=1.0)
    hr = hp.add_run("IESPFLIX  |  Guia de Estudo Atualizado")
    set_font(hr, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(fp, after=0, line=1.0)
    fr = fp.add_run("Projeto Flix  •  Samantha Hellen  •  Página ")
    set_font(fr, size=8.5, color=MUTED)
    add_page_number(fp)


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    bullet_num, decimal_num = setup_numbering(doc)

    # Cover
    add_para(doc, "GUIA DE ESTUDO", bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=72, after=16)
    p = add_para(doc, "IESPFLIX", bold=True, color=INK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    set_font(p.runs[0], size=30, bold=True, color=INK_BLUE)
    p = add_para(doc, "Projeto Flix Backend", color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    set_font(p.runs[0], size=17, color=DARK_BLUE)
    add_para(
        doc,
        "Spring Boot, requisitos obrigatórios, código atual e roteiro para apresentação",
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=36,
    )
    add_callout(
        doc,
        "Versão estudada",
        "Guia reconstruído a partir do código atual validado em 01/06/2026. "
        "A suíte executada possui 55 testes sem falhas. Este material substitui os guias antigos.",
        fill=LIGHT_BLUE,
    )
    add_para(doc, "Aluno(a): Samantha Hellen", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=20, after=4)
    add_para(doc, "Disciplina: Tecnologias para Backend", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Projeto: API REST de streaming IESPFLIX", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Professor: Rodrigo Fujioka", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Data de atualização: 01/06/2026", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_page_break(doc)

    # Contents and how to use
    add_heading(doc, "Como usar este guia", 1)
    add_para(
        doc,
        "Este documento foi escrito para duas situações: estudar o projeto do zero e revisar rapidamente "
        "antes da apresentação. Leia primeiro as Partes I e II para entender Spring Boot e a organização "
        "do código. Depois estude a Parte III, que liga cada requisito obrigatório aos arquivos atuais. "
        "Na véspera da apresentação, use as Partes V e VI como roteiro prático.",
    )
    add_callout(
        doc,
        "Resumo em uma frase",
        "O IESPFLIX é uma API REST em Spring Boot para cadastro de usuários, autenticação, catálogo de "
        "filmes e séries, favoritos e cartões tokenizados, com persistência H2, validações, Swagger, "
        "Spring Security e integração externa ViaCEP.",
    )
    add_heading(doc, "Sumário", 2)
    for item in [
        "Parte I - Fundamentos: API REST, Spring Boot, Maven e injeção de dependência",
        "Parte II - Anatomia do projeto atual: pacotes, banco, DTOs, validações e fluxo",
        "Parte III - Requisitos obrigatórios: RF1, RF3, RF4, RF7, RF8, RF10, RF11, RNF1 e RNF2",
        "Parte IV - Tópicos avaliados: JPQL, padrões, tratamento de erros, ViaCEP, testes e JaCoCo",
        "Parte V - Demonstração no Swagger e inspeção pelo H2",
        "Parte VI - Roteiro oral, perguntas prováveis, checklist e glossário",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "O que mudou em relação aos guias antigos", 2)
    add_table(
        doc,
        ["Antes", "Agora no código atual", "Motivo"],
        [
            ("Java 21", "Java 17", "Compatibilidade com o ambiente instalado e compilação reproduzível."),
            ("Cadastro recebia UsuarioDTO", "Cadastro recebe CadastroUsuarioDTO", "RF1 exige confirmação de senha e dados do cartão no mesmo cadastro."),
            ("Cartão explicado como CRUD simples", "Número completo e CVV não são persistidos", "Reduz exposição de dados sensíveis; ficam token, bandeira e últimos 4 dígitos."),
            ("Admin descrito sem Swagger visual", "OpenApiConfig cria basicAuth e botão Authorize", "Facilita demonstrar RF11 pelo Swagger."),
            ("Tipo de conteúdo aberto", "ConteudoDTO aceita somente FILME ou SERIE", "Evita registros inválidos como STRING."),
            ("Filme era o foco", "Conteudo é o catálogo principal", "Um único modelo representa filmes e séries e permite separar favoritos."),
        ],
        [1800, 3300, 4260],
        font_size=8.8,
    )

    add_page_break(doc)

    # Part I
    add_heading(doc, "Parte I - Fundamentos de Spring Boot", 1)
    add_heading(doc, "1. O que é backend?", 2)
    add_para(
        doc,
        "Backend é a parte do sistema executada no servidor. Ele recebe requisições, aplica regras de "
        "negócio, consulta ou altera dados e devolve respostas. No IESPFLIX, o navegador ou Swagger envia "
        "JSON para a API, e o backend responde com JSON e um código HTTP.",
    )
    add_heading(doc, "2. O que é uma API REST?", 2)
    add_para(
        doc,
        "API é uma interface para comunicação entre sistemas. REST é um estilo arquitetural que organiza "
        "recursos por URLs e usa os verbos do HTTP. No projeto, usuários, conteúdos, favoritos e cartões "
        "são recursos acessíveis por endpoints.",
    )
    add_table(
        doc,
        ["Verbo", "Uso comum", "Exemplo no projeto", "Resposta típica"],
        [
            ("GET", "Consultar", "GET /api/v1/conteudos", "200 OK"),
            ("POST", "Criar", "POST /api/v1/usuarios", "201 Created"),
            ("PUT", "Atualizar por completo", "PUT /api/v1/metodos-pagamento/{id}", "200 OK"),
            ("PATCH", "Alterar parcialmente", "PATCH /api/v1/assinaturas/{id}/cancelar", "200 OK"),
            ("DELETE", "Excluir", "DELETE /api/v1/conteudos/{id}", "204 No Content"),
        ],
        [900, 1800, 4260, 2400],
    )
    add_heading(doc, "3. Códigos HTTP que você deve saber", 2)
    add_table(
        doc,
        ["Código", "Significado", "Quando aparece"],
        [
            ("200", "OK", "Consulta, login ou atualização concluída."),
            ("201", "Created", "Cadastro criado, como usuário, conteúdo ou favorito."),
            ("204", "No Content", "Exclusão concluída sem corpo na resposta."),
            ("400", "Bad Request", "JSON inválido ou falha de Bean Validation."),
            ("401", "Unauthorized", "Credenciais ausentes ou incorretas."),
            ("403", "Forbidden", "Usuário autenticado, mas sem permissão de administrador."),
            ("404", "Not Found", "Recurso não encontrado."),
            ("409", "Conflict", "E-mail, CPF/CNPJ ou favorito duplicado."),
            ("500", "Internal Server Error", "Falha inesperada não tratada especificamente."),
        ],
        [900, 2400, 6060],
    )
    add_heading(doc, "4. O que é Spring Boot?", 2)
    add_para(
        doc,
        "Spring Boot é um framework Java que facilita criar aplicações web. Ele configura automaticamente "
        "o servidor, integra bibliotecas e encontra classes anotadas. Em vez de montar toda a infraestrutura "
        "manualmente, você declara intenções com anotações.",
    )
    add_table(
        doc,
        ["Anotação", "O que informa ao Spring", "Onde aparece"],
        [
            ("@SpringBootApplication", "Classe principal e configuração automática.", "TechbackApplication.java"),
            ("@RestController", "Classe que recebe requisições HTTP e devolve JSON.", "Controllers"),
            ("@RequestMapping", "Prefixo da URL de um controller.", "Ex.: /api/v1/conteudos"),
            ("@Service", "Classe com regra de negócio.", "Services"),
            ("@Repository", "Componente de acesso ao banco.", "Repositories"),
            ("@Entity", "Classe persistida como tabela.", "Models"),
            ("@Configuration", "Classe que fornece configurações e beans.", "config"),
            ("@Bean", "Objeto criado e gerenciado pelo Spring.", "SecurityConfig, ModelMapperConfig"),
        ],
        [1900, 4000, 3460],
        font_size=9.0,
    )
    add_heading(doc, "5. Classe principal e inicialização", 2)
    add_para(doc, "A aplicação começa em src/main/java/br/uniesp/si/techback/TechbackApplication.java.")
    add_code(
        doc,
        """@SpringBootApplication
@EnableFeignClients
public class TechbackApplication {
    public static void main(String[] args) {
        SpringApplication.run(TechbackApplication.class, args);
    }
}""",
    )
    add_para(
        doc,
        "@SpringBootApplication inicia o projeto e procura componentes nos subpacotes. "
        "@EnableFeignClients habilita interfaces declarativas para chamar APIs externas, como o ViaCEP.",
    )
    add_heading(doc, "6. Maven e pom.xml", 2)
    add_para(
        doc,
        "O Maven baixa bibliotecas, compila e executa testes. O arquivo pom.xml é a lista de dependências "
        "e plugins do projeto. A propriedade java.version está configurada como 17.",
    )
    add_table(
        doc,
        ["Dependência ou plugin", "Função no projeto"],
        [
            ("spring-boot-starter-web", "API REST, JSON e servidor web embutido."),
            ("spring-boot-starter-data-jpa", "Persistência com JPA e Hibernate."),
            ("h2", "Banco local leve para desenvolvimento."),
            ("spring-boot-starter-validation", "Bean Validation: @NotBlank, @Pattern, @Email etc."),
            ("spring-boot-starter-security", "HTTP Basic, autorização por perfil e BCrypt."),
            ("spring-cloud-starter-openfeign", "Cliente HTTP declarativo para ViaCEP."),
            ("springdoc-openapi-starter-webmvc-ui", "Documentação Swagger UI."),
            ("flyway-core", "Suporte a migrações; desativado no perfil dev atual."),
            ("lombok", "Gera getters, setters, builders e construtores."),
            ("spring-boot-starter-test", "JUnit, Mockito, MockMvc e ferramentas de teste."),
            ("jacoco-maven-plugin", "Relatório de cobertura em target/site/jacoco/index.html."),
        ],
        [3300, 6060],
    )
    add_callout(
        doc,
        "Comandos essenciais",
        "Iniciar: .\\mvnw.cmd spring-boot:run   |   Testar: .\\mvnw.cmd test   |   Swagger: "
        "http://localhost:8080/swagger-ui/index.html",
        fill=LIGHT_BLUE,
    )

    add_page_break(doc)

    # Part II
    add_heading(doc, "Parte II - Anatomia do projeto atual", 1)
    add_heading(doc, "7. Arquitetura em camadas", 2)
    add_para(
        doc,
        "O projeto aplica separação de responsabilidades. Cada camada resolve um tipo de problema. "
        "Essa organização facilita manutenção, teste e explicação durante a apresentação.",
    )
    add_table(
        doc,
        ["Camada", "Responsabilidade", "Exemplo atual"],
        [
            ("Controller", "Recebe HTTP, lê parâmetros e chama service.", "UsuarioController, ConteudoController"),
            ("DTO", "Define dados de entrada e saída da API.", "CadastroUsuarioDTO, ConteudoDTO"),
            ("Validation", "Bloqueia dados inválidos antes da regra de negócio.", "CpfCnpjValidator, SenhaForteValidator"),
            ("Service", "Aplica regras e coordena operações.", "UsuarioService, MetodoPagamentoService"),
            ("Repository", "Acessa o banco com Spring Data JPA.", "UsuarioRepository, ConteudoRepository"),
            ("Model / Entity", "Representa tabelas do banco.", "Usuario, Conteudo, Favorito"),
            ("Config", "Define segurança, Swagger e beans.", "SecurityConfig, OpenApiConfig"),
            ("Exception", "Transforma erros em respostas padronizadas.", "GlobalExceptionHandler"),
        ],
        [1500, 4300, 3560],
        font_size=9.0,
    )
    add_heading(doc, "8. Fluxo completo de uma requisição", 2)
    for text in [
        "O Swagger envia uma requisição HTTP com JSON.",
        "O Controller recebe a requisição e usa @Valid no DTO.",
        "O Bean Validation verifica formato, obrigatoriedade e validadores customizados.",
        "O Service aplica regras de negócio, como duplicidade, autorização e hash.",
        "O Repository consulta ou salva a Entity no H2.",
        "O Service monta um DTO de resposta sem expor segredos.",
        "O Controller devolve JSON e um código HTTP.",
        "Se houver erro, o GlobalExceptionHandler padroniza a resposta.",
    ]:
        add_number(doc, text, decimal_num)
    add_callout(
        doc,
        "Fluxo resumido",
        "Swagger -> Controller -> DTO + Validation -> Service -> Repository -> H2 -> DTO de resposta -> JSON",
    )

    add_heading(doc, "9. Pacotes do núcleo IESPFLIX", 2)
    add_table(
        doc,
        ["Pacote", "Classes principais", "Por que estudar"],
        [
            ("controller", "UsuarioController, AuthController, ConteudoController, FavoritoController, OutrosControllers", "Endpoints centrais da demonstração."),
            ("service", "UsuarioService, AuthService, ConteudoService, FavoritoService, MetodoPagamentoService", "Regras obrigatórias."),
            ("repository", "UsuarioRepository, ConteudoRepository, FavoritoRepository, MetodoPagamentoRepository", "Persistência e JPQL."),
            ("model", "Usuario, MetodoPagamento, Conteudo, Favorito, Plano, Assinatura", "Tabelas do domínio."),
            ("dto", "CadastroUsuarioDTO, AtualizacaoCartaoDTO, ConteudoDTO, LoginRequestDTO", "Contrato JSON e validações."),
            ("config", "SecurityConfig, OpenApiConfig", "Autorização e botão Authorize."),
            ("validation", "CpfCnpj, SenhaForte e validadores", "Bean Validation customizada."),
        ],
        [1500, 4400, 3460],
        font_size=8.8,
    )
    add_heading(doc, "10. Módulos didáticos legados", 2)
    add_para(
        doc,
        "O repositório também contém módulos criados anteriormente para praticar CRUD, mapper, JPQL e serviço "
        "externo: Filme, Produto, Cliente, Categoria, Pedido e Funcionario. Eles continuam úteis, mas não são o "
        "catálogo principal da apresentação.",
    )
    add_callout(
        doc,
        "Não confunda",
        "Para RF8, RF10 e RF11 use Conteudo e as rotas /api/v1/conteudos. A entidade Filme e a rota /filmes "
        "são exemplos antigos preservados no projeto.",
        fill="FFF4D6",
        title_color=GOLD,
    )

    add_heading(doc, "11. JPA, Hibernate e entidades", 2)
    add_para(
        doc,
        "JPA é uma especificação Java para persistência. Hibernate é a implementação usada pelo Spring Boot. "
        "Uma classe com @Entity vira uma tabela. @Id indica chave primária e @GeneratedValue deixa o banco gerar o ID.",
    )
    add_table(
        doc,
        ["Entidade", "Tabela", "Campos importantes", "Uso"],
        [
            ("Usuario", "usuarios", "email, senhaHash, cpfCnpj, perfil", "Cadastro, login e autorização."),
            ("MetodoPagamento", "metodo_pagamento", "ultimos4, bandeira, tokenGateway", "Cartão sem armazenar número completo ou CVV."),
            ("Conteudo", "conteudo", "titulo, tipo, ano, duração, relevância, sinopse, trailerUrl, gênero", "Catálogo de filmes e séries."),
            ("Favorito", "favoritos", "usuarioId, conteudoId", "Relaciona usuário ao conteúdo favorito."),
            ("Plano", "plano", "codigo, limiteDiario, preco", "Bonificação: tipos de plano."),
            ("Assinatura", "assinatura", "usuarioId, planoId, status, datas", "Bonificação: vínculo do usuário com plano."),
        ],
        [1450, 1600, 3900, 2410],
        font_size=8.6,
    )
    add_code(
        doc,
        """@Entity
@Table(name = "conteudo")
public class Conteudo {
    @Id
    @GeneratedValue(strategy = GenerationType.IDENTITY)
    private Long id;
    private String titulo;
    private String tipo; // FILME ou SERIE
    private Integer ano;
    private Integer duracaoMinutos;
    private BigDecimal relevancia;
    private String sinopse;
    private String trailerUrl;
    private String genero;
}""",
    )

    add_heading(doc, "12. DTOs: contrato da API", 2)
    add_para(
        doc,
        "DTO significa Data Transfer Object. DTOs controlam o que entra e sai da API. Eles evitam expor a "
        "estrutura inteira da entidade e são o lugar ideal para Bean Validation.",
    )
    add_table(
        doc,
        ["DTO", "Uso"],
        [
            ("CadastroUsuarioDTO", "Entrada completa do RF1: dados pessoais, senha confirmada e cartão."),
            ("UsuarioDTO", "Resposta e edição de usuário sem devolver senha hash."),
            ("LoginRequestDTO / LoginResponseDTO", "Entrada e saída da autenticação."),
            ("AtualizacaoCartaoDTO", "Entrada do RF7 com número, validade, CVV e titular."),
            ("MetodoPagamentoDTO", "Representação persistida e operações adicionais de cartão."),
            ("ConteudoDTO", "Contrato do catálogo com validação de FILME ou SERIE."),
            ("FavoritoDTO", "Relaciona usuarioId e conteudoId."),
        ],
        [2900, 6460],
    )
    add_callout(
        doc,
        "Segurança de saída",
        "@JsonProperty(access = WRITE_ONLY) permite receber dados sensíveis no JSON de entrada, mas impede "
        "que eles apareçam na resposta. O projeto aplica isso à senha, confirmação de senha, número do cartão, "
        "CVV, validade recebida e token do gateway.",
    )

    add_heading(doc, "13. Bean Validation", 2)
    add_para(
        doc,
        "Bean Validation bloqueia entradas erradas antes de o service executar. O Controller ativa isso com "
        "@Valid. O projeto combina anotações prontas e validadores próprios.",
    )
    add_table(
        doc,
        ["Anotação", "Validação", "Exemplo"],
        [
            ("@NotBlank", "Texto obrigatório e não vazio.", "nomeCompleto, email, senha"),
            ("@NotNull", "Valor obrigatório.", "dataNascimento, ano, duração"),
            ("@Email", "Formato básico de e-mail.", "CadastroUsuarioDTO.email"),
            ("@Past", "Data deve estar no passado.", "dataNascimento"),
            ("@Pattern", "Texto precisa casar com regex.", "cartão, validade, CVV, tipo"),
            ("@Min / @Max", "Intervalo numérico.", "ano, duração, relevância"),
            ("@CpfCnpj", "Dígitos verificadores de CPF ou CNPJ.", "cpfCnpj"),
            ("@SenhaForte", "8+ caracteres, maiúscula, minúscula, número e especial.", "senha"),
        ],
        [1800, 3900, 3660],
    )
    add_code(
        doc,
        """@NotBlank(message = "Tipo é obrigatório (FILME ou SERIE)")
@Pattern(regexp = "(?i)FILME|SERIE",
         message = "Tipo deve ser FILME ou SERIE")
private String tipo;""",
    )

    add_heading(doc, "14. H2 e persistência", 2)
    add_para(
        doc,
        "O perfil dev usa H2 em arquivo. Isso significa que os dados continuam salvos após fechar e abrir o "
        "projeto. O perfil test usa H2 em memória e create-drop, portanto os testes não sujam o banco usado na apresentação.",
    )
    add_table(
        doc,
        ["Ambiente", "URL JDBC", "DDL auto", "Comportamento"],
        [
            ("dev", "jdbc:h2:file:~/teckback20262", "update", "Mantém dados entre reinícios."),
            ("test", "jdbc:h2:mem:testdb", "create-drop", "Cria banco temporário e apaga ao terminar."),
        ],
        [1200, 3600, 1500, 3060],
    )
    add_para(doc, "Console H2: http://localhost:8080/h2")
    add_code(
        doc,
        """JDBC URL: jdbc:h2:file:~/teckback20262
User Name: sa
Password:  [deixe vazio]""",
    )

    # Part III
    add_heading(doc, "Parte III - Requisitos do professor no código atual", 1)
    add_heading(doc, "15. Mapa geral dos requisitos obrigatórios", 2)
    add_table(
        doc,
        ["ID", "Pedido do professor", "Implementação atual", "Endpoint principal"],
        [
            ("RF1", "Cadastrar usuário com dados pessoais, senha confirmada e cartão.", "CadastroUsuarioDTO + UsuarioService + MetodoPagamentoService", "POST /api/v1/usuarios"),
            ("RF3", "Salvar senha com HASH.", "BCryptPasswordEncoder e campo senhaHash", "POST /api/v1/usuarios"),
            ("RF4", "API de login.", "AuthController + AuthService", "POST /api/v1/auth/login"),
            ("RF7", "Alterar cartão após autenticação.", "AtualizacaoCartaoDTO + HTTP Basic + autorização por dono", "PUT /api/v1/metodos-pagamento/{id}"),
            ("RF8", "Detalhar filmes e séries.", "Conteudo + ConteudoDTO", "GET /api/v1/conteudos/{id}"),
            ("RF10", "Listar favoritos por tipo.", "FavoritoRepository filtra FILME ou SERIE", "GET .../filmes e GET .../series"),
            ("RF11", "Exigir admin para cadastrar filmes e séries.", "SecurityConfig hasRole(ADMIN)", "POST /api/v1/conteudos"),
            ("RNF1", "Versionamento Git.", "Repositório Git com remote origin.", "GitHub do projeto"),
            ("RNF2", "README executável.", "README.md com execução, Swagger, H2 e fluxos.", "README.md"),
        ],
        [680, 2600, 3550, 2530],
        font_size=8.2,
    )

    add_heading(doc, "16. RF1 - Cadastro completo do usuário", 2)
    add_para(
        doc,
        "O cadastro público fica em POST /api/v1/usuarios. O Controller recebe CadastroUsuarioDTO, "
        "não UsuarioDTO. Essa separação existe porque o cadastro inicial precisa de campos extras: "
        "confirmarSenha e os dados do cartão.",
    )
    add_code(
        doc,
        """@PostMapping
@ResponseStatus(HttpStatus.CREATED)
public UsuarioDTO criar(@Valid @RequestBody CadastroUsuarioDTO dto) {
    return usuarioService.criar(dto);
}""",
    )
    add_para(doc, "Dentro de UsuarioService.criar, a ordem principal é:")
    for text in [
        "Verificar se o e-mail já existe.",
        "Verificar se CPF/CNPJ já existe.",
        "Comparar senha e confirmação.",
        "Forçar perfil USER, impedindo cadastro público como ADMIN.",
        "Gerar hash BCrypt e salvar Usuario.",
        "Cadastrar cartão inicial tokenizado dentro da mesma transação.",
        "Devolver UsuarioDTO sem segredos.",
    ]:
        add_number(doc, text, decimal_num)
    add_callout(
        doc,
        "Por que @Transactional?",
        "Se salvar o usuário funcionar, mas cadastrar o cartão falhar, a transação evita deixar um cadastro "
        "incompleto. A operação inteira é confirmada ou desfeita.",
    )

    add_heading(doc, "17. RF3 - Senha com hash BCrypt", 2)
    add_para(
        doc,
        "O banco não guarda a senha digitada. UsuarioService chama passwordEncoder.encode e salva o resultado "
        "no campo senhaHash. O BCrypt inclui salt e gera hashes diferentes mesmo para senhas iguais.",
    )
    add_code(doc, 'usuario.setSenhaHash(passwordEncoder.encode(dto.getSenha()));')
    add_para(
        doc,
        "Durante o login, a comparação não é feita com equals. AuthService usa passwordEncoder.matches para "
        "comparar a senha digitada ao hash persistido.",
    )
    add_code(doc, "boolean senhaConfere = passwordEncoder.matches(dto.getSenha(), usuario.getSenhaHash());")
    add_callout(
        doc,
        "Frase pronta",
        "A senha nunca é persistida em texto puro. No cadastro eu salvo o hash BCrypt em senhaHash e, no login, "
        "uso matches para validar a senha informada.",
    )

    add_heading(doc, "18. RF4 - API de login", 2)
    add_para(doc, "A rota é POST /api/v1/auth/login. Ela recebe LoginRequestDTO e devolve LoginResponseDTO.")
    add_code(
        doc,
        """{
  "email": "maria@example.com",
  "senha": "Senha@123"
}""",
    )
    add_para(
        doc,
        "AuthService busca o usuário pelo e-mail. Se não achar ou a senha estiver errada, retorna 401. "
        "Se estiver certa, devolve id, nome, e-mail, perfil e mensagem de sucesso.",
    )
    add_callout(
        doc,
        "Atenção: login x rotas protegidas",
        "O endpoint /auth/login comprova as credenciais, mas não emite JWT nem cria sessão. Para acessar rotas "
        "protegidas, o projeto usa HTTP Basic: o cliente envia usuário e senha a cada requisição.",
        fill="FFF4D6",
        title_color=GOLD,
    )

    add_heading(doc, "19. RF7 - Alteração do cartão após autenticação", 2)
    add_para(
        doc,
        "As rotas /api/v1/metodos-pagamento/** exigem autenticação. O PUT recebe AtualizacaoCartaoDTO e "
        "Authentication. O Spring Security informa o e-mail autenticado por authentication.getName().",
    )
    add_code(
        doc,
        """@PutMapping("/{id}")
public MetodoPagamentoDTO atualizar(
        @PathVariable Long id,
        @Valid @RequestBody AtualizacaoCartaoDTO dto,
        Authentication authentication) {
    return metodoPagamentoService.atualizar(
            id, dto, authentication.getName());
}""",
    )
    add_para(doc, "O service valida o dono do cartão e armazena somente dados reduzidos:")
    add_table(
        doc,
        ["Recebido na entrada", "Persistido no banco?", "O que fica salvo"],
        [
            ("Número completo", "Não", "Somente os últimos 4 dígitos e bandeira identificada."),
            ("CVV", "Não", "Nunca é persistido."),
            ("Validade", "Sim, parcialmente", "Mês e ano."),
            ("Titular", "Sim", "Nome do portador."),
            ("Token", "Gerado localmente", "UUID no campo tokenGateway."),
        ],
        [3000, 1800, 4560],
    )
    add_para(
        doc,
        "MetodoPagamentoService.autorizarUsuario permite o próprio usuário e também o administrador. Outro "
        "usuário autenticado recebe 403 Forbidden.",
    )

    add_heading(doc, "20. RF8 - Catálogo detalhado de filmes e séries", 2)
    add_para(
        doc,
        "Conteudo representa filmes e séries em uma tabela única. O campo tipo diferencia FILME e SERIE. "
        "Isso evita duplicar lógica e também facilita separar favoritos.",
    )
    add_table(
        doc,
        ["Campo RF8", "Campo Java", "Exemplo"],
        [
            ("Título", "titulo", "Coringa"),
            ("Gênero", "genero", "Drama"),
            ("Ano", "ano", "2019"),
            ("Duração", "duracaoMinutos", "122"),
            ("Relevância", "relevancia", "8.5"),
            ("Sinopse", "sinopse", "A origem de um vilão conhecido..."),
            ("Trailer", "trailerUrl", "https://youtu.be/..."),
        ],
        [1800, 2600, 4960],
    )
    add_para(doc, "Endpoints úteis:")
    for text in [
        "GET /api/v1/conteudos - lista tudo em ordem alfabética.",
        "GET /api/v1/conteudos/{id} - mostra todos os detalhes.",
        "GET /api/v1/conteudos?tipo=FILME - filtra filmes.",
        "GET /api/v1/conteudos?tipo=SERIE - filtra séries.",
        "GET /api/v1/conteudos?genero=Drama - filtra por gênero.",
        "GET /api/v1/conteudos?q=Dark - pesquisa no título.",
        "GET /api/v1/conteudos/top-relevancia - ordena por relevância.",
        "GET /api/v1/conteudos/lancados-apos?ano=2015 - filtra por ano.",
    ]:
        add_bullet(doc, text, bullet_num)

    add_heading(doc, "21. RF10 e RF9 - Favoritos", 2)
    add_para(
        doc,
        "RF10 exige listar filmes e séries separadamente. O projeto também implementa RF9, embora ele não seja "
        "obrigatório: é possível adicionar um favorito por usuarioId e conteudoId.",
    )
    add_code(
        doc,
        """POST /api/v1/favoritos
{
  "usuarioId": 1,
  "conteudoId": 41
}""",
    )
    add_para(doc, "FavoritoService verifica três condições antes de salvar:")
    for text in [
        "O usuário existe.",
        "O conteúdo existe.",
        "A mesma combinação usuário + conteúdo ainda não foi favoritada.",
    ]:
        add_number(doc, text, decimal_num)
    add_para(doc, "Para listar separadamente:")
    add_code(
        doc,
        """GET /api/v1/favoritos/usuario/{usuarioId}/filmes
GET /api/v1/favoritos/usuario/{usuarioId}/series""",
    )
    add_para(
        doc,
        "FavoritoRepository usa JPQL com JOIN entre Favorito e Conteudo e filtra UPPER(c.tipo). "
        "Assim, um endpoint devolve apenas FILME e o outro apenas SERIE.",
    )

    add_heading(doc, "22. RF11 - Administrador para cadastrar conteúdos", 2)
    add_para(
        doc,
        "SecurityConfig protege POST, PUT e DELETE de /api/v1/conteudos com hasRole(\"ADMIN\"). "
        "Consultas GET continuam públicas.",
    )
    add_code(
        doc,
        """.requestMatchers(HttpMethod.POST,
        "/api/v1/conteudos", "/api/v1/conteudos/**")
    .hasRole("ADMIN")""",
    )
    add_para(doc, "Credenciais locais de demonstração:")
    add_code(
        doc,
        """Username: admin
Password: admin123""",
    )
    add_para(
        doc,
        "OpenApiConfig registra o esquema basicAuth para o Swagger exibir o botão Authorize. Depois de autorizar, "
        "o Swagger envia o cabeçalho Authorization: Basic ... automaticamente.",
    )
    add_callout(
        doc,
        "Detalhe técnico importante",
        "SecurityConfig cria um novo objeto User para o administrador a cada autenticação. Isso evita que o "
        "apagamento interno de credenciais pelo Spring inutilize o admin após o primeiro uso.",
    )

    add_heading(doc, "23. RNF1 e RNF2", 2)
    add_table(
        doc,
        ["Requisito", "Como está atendido", "Como demonstrar"],
        [
            ("RNF1 - Git", "Projeto versionado no repositório Projeto_TecBack.", "Mostrar histórico e remote origin no GitHub."),
            ("RNF2 - README", "README.md contém pré-requisitos, execução, H2, Swagger e exemplos.", "Abrir README.md e executar o Maven Wrapper."),
        ],
        [1900, 4160, 3300],
    )

    # Part IV
    add_heading(doc, "Parte IV - Tópicos cobrados na avaliação", 1)
    add_heading(doc, "24. Critérios descritos pelo professor", 2)
    add_table(
        doc,
        ["Critério", "Peso", "Onde aparece no projeto"],
        [
            ("Entregas na data certa", "20%", "Processo acadêmico; não depende do código."),
            ("Bean Validation", "20%", "DTOs, @Valid, @CpfCnpj, @SenhaForte e @Pattern."),
            ("Services e serviço externo", "20%", "Camada service e integração ViaCEP com OpenFeign."),
            ("Persistência", "10%", "JPA, Hibernate, repositories e H2 em arquivo."),
            ("Padrões de projeto", "10%", "Camadas, Repository, DTO, Builder, DI e handler global."),
            ("Git", "10%", "Controle de versão e GitHub."),
            ("Swagger", "10%", "SpringDoc OpenAPI, Swagger UI e basicAuth."),
        ],
        [3000, 900, 5460],
    )

    add_heading(doc, "25. JPQL no ConteudoRepository", 2)
    add_para(
        doc,
        "JPQL se parece com SQL, mas consulta classes e atributos Java em vez de nomes físicos de tabelas. "
        "ConteudoRepository possui cinco consultas com @Query e um método derivado.",
    )
    add_table(
        doc,
        ["Consulta", "Trecho principal", "Uso"],
        [
            ("JPQL 1", "SELECT c FROM Conteudo c ORDER BY c.titulo ASC", "Lista alfabética."),
            ("JPQL 2", "WHERE LOWER(c.genero) = LOWER(:genero)", "Filtro sem diferenciar maiúsculas."),
            ("JPQL 3", "ORDER BY c.relevancia DESC", "Ranking de relevância."),
            ("JPQL 4", "WHERE c.ano > :ano ORDER BY c.ano DESC", "Lançados após um ano."),
            ("JPQL 5", "LOWER(c.titulo) LIKE LOWER(CONCAT('%', :q, '%'))", "Pesquisa por palavra no título."),
            ("Derived Query", "findAllByTipoOrderByTituloAsc(tipo)", "Filtro por FILME ou SERIE."),
        ],
        [1100, 5030, 3230],
        font_size=8.6,
    )
    add_callout(
        doc,
        "Frase pronta",
        "Em JPQL eu uso o nome da entidade Conteudo e seus atributos Java. O Hibernate traduz para SQL e consulta "
        "a tabela física no H2.",
    )

    add_heading(doc, "26. Padrões e boas práticas presentes", 2)
    add_table(
        doc,
        ["Prática ou padrão", "Exemplo no projeto", "Benefício"],
        [
            ("Arquitetura em camadas", "Controller -> Service -> Repository", "Responsabilidades separadas."),
            ("Repository Pattern", "Interfaces que estendem JpaRepository", "Centraliza acesso a dados."),
            ("DTO Pattern", "CadastroUsuarioDTO, ConteudoDTO", "Controla contrato e evita exposição indevida."),
            ("Dependency Injection", "@RequiredArgsConstructor com campos final", "Baixo acoplamento e testabilidade."),
            ("Builder Pattern", "Lombok @Builder", "Criação legível de objetos."),
            ("Declarative HTTP Client", "ViaCepClient com @FeignClient", "Integração externa simples."),
            ("Global Exception Handler", "@RestControllerAdvice", "Erros consistentes."),
            ("Tokenização demonstrativa", "UUID em tokenGateway", "Não persiste número completo nem CVV."),
        ],
        [2800, 3300, 3260],
        font_size=8.8,
    )

    add_heading(doc, "27. Tratamento global de erros", 2)
    add_para(
        doc,
        "GlobalExceptionHandler centraliza respostas de erro e usa ProblemDetail. Isso evita repetir tratamento "
        "em cada controller e mantém status corretos.",
    )
    add_table(
        doc,
        ["Exceção", "Resposta", "Exemplo"],
        [
            ("MethodArgumentNotValidException", "400 + errors por campo", "Senha fraca ou tipo STRING."),
            ("CustomBeanException", "400", "CEP inválido retornado pelo ViaCEP."),
            ("ResponseStatusException", "Status original", "401, 403, 404 ou 409 criados pelos services."),
            ("Exception genérica", "500", "Erro inesperado."),
        ],
        [3100, 2600, 3660],
    )
    add_para(
        doc,
        "O Content-Type da resposta é application/problem+json. Essa estrutura segue o formato ProblemDetail "
        "do Spring para comunicar erros de forma padronizada.",
    )

    add_heading(doc, "28. Serviço externo ViaCEP", 2)
    add_para(
        doc,
        "O professor pontua uso de service e serviço externo. O projeto demonstra isso no módulo Funcionario. "
        "Quando um funcionário é cadastrado com CEP, FuncionarioService remove caracteres, chama ViaCepClient e "
        "preenche logradouro, bairro, localidade e UF.",
    )
    add_code(
        doc,
        """@FeignClient(name = "viaCepClient",
             url = "${viacep.url:https://viacep.com.br/ws}")
public interface ViaCepClient {
    @GetMapping("/{cep}/json/")
    ViaCepResponseDTO buscarPorCep(@PathVariable("cep") String cep);
}""",
    )
    add_para(doc, "Endpoint de demonstração:")
    add_code(
        doc,
        """POST /funcionarios
{
  "nome": "Samantha",
  "cargo": "Analista",
  "cep": "58000-000"
}""",
    )
    add_callout(
        doc,
        "Contexto",
        "Funcionario é um módulo didático complementar, separado do domínio principal de streaming. Ele existe "
        "para demonstrar integração externa com OpenFeign.",
    )

    add_heading(doc, "29. Swagger e OpenAPI", 2)
    add_para(
        doc,
        "Swagger UI é a documentação interativa gerada pelo SpringDoc. Ele lista endpoints, mostra schemas JSON "
        "e permite executar requisições. OpenApiConfig adiciona basicAuth para que o botão Authorize apareça.",
    )
    add_code(
        doc,
        """Swagger UI:  http://localhost:8080/swagger-ui/index.html
OpenAPI JSON: http://localhost:8080/v3/api-docs""",
    )

    add_heading(doc, "30. Testes automatizados e JaCoCo", 2)
    add_para(
        doc,
        "Os testes servem como prova repetível de comportamento. A execução final deste guia passou com 55 testes, "
        "0 falhas e 0 erros. O relatório JaCoCo foi gerado com cobertura total de 45% das instruções e 47% dos branches.",
    )
    add_table(
        doc,
        ["Tipo de teste", "Arquivo", "O que comprova"],
        [
            ("Integração", "ObrigatoriosIntegrationTest", "Fluxo obrigatório completo, HTTP Basic e respostas HTTP."),
            ("Repository", "ConteudoRepositoryTest", "Quatro consultas JPQL principais."),
            ("Service", "UsuarioServiceTest", "Hash, perfil USER e cartão inicial."),
            ("Service", "MetodoPagamentoServiceTest", "Tokenização e bloqueio de outro usuário."),
            ("Validation", "CpfCnpjValidatorTest", "CPF/CNPJ válidos e inválidos."),
            ("Validation", "SenhaForteValidatorTest", "Regras de senha forte."),
            ("Legado", "FilmeControllerTest, FilmeServiceTest etc.", "CRUD e estrutura didática original."),
        ],
        [1700, 3100, 4560],
        font_size=8.8,
    )
    add_code(
        doc,
        """.\\mvnw.cmd test

# relatório após os testes
target/site/jacoco/index.html""",
    )

    add_heading(doc, "31. Bonificações e itens não obrigatórios", 2)
    add_table(
        doc,
        ["Item", "Situação atual", "Observação honesta para apresentar"],
        [
            ("RF2 - confirmação por e-mail", "Não implementado", "Não é obrigatório."),
            ("RF9 - adicionar favorito", "Implementado", "POST /api/v1/favoritos."),
            ("RFB1 - planos diferentes", "Estrutura implementada", "Plano e Assinatura existem com endpoints."),
            ("RFB2 - limitar instâncias simultâneas", "Não implementado", "Não é obrigatório."),
            ("Flyway", "Dependência presente", "Desativado no perfil dev atual; testes possuem migration."),
            ("Testcontainers", "Dependências presentes", "Não é o foco da demonstração atual."),
        ],
        [2800, 1900, 4660],
        font_size=8.8,
    )

    # Part V
    add_heading(doc, "Parte V - Demonstração prática", 1)
    add_heading(doc, "32. Preparar o ambiente", 2)
    for text in [
        "Abra um terminal na pasta Projeto Flix.",
        "Execute .\\mvnw.cmd spring-boot:run.",
        "Aguarde o servidor iniciar na porta 8080.",
        "Abra http://localhost:8080/swagger-ui/index.html.",
        "Use Ctrl+F5 se o navegador estiver mostrando uma versão antiga.",
    ]:
        add_number(doc, text, decimal_num)
    add_callout(
        doc,
        "Se aparecer um pop-up do navegador pedindo login",
        "Clique em Cancelar. Esse pop-up costuma ser uma resposta 401. Para a demonstração normal, use o botão "
        "Authorize dentro do Swagger e informe as credenciais adequadas.",
        fill="FFF4D6",
        title_color=GOLD,
    )

    add_heading(doc, "33. Roteiro recomendado no Swagger", 2)
    add_para(doc, "O roteiro abaixo demonstra os principais requisitos em uma ordem fácil de explicar.")
    for text in [
        "Cadastrar um usuário em POST /api/v1/usuarios.",
        "Fazer login em POST /api/v1/auth/login.",
        "Abrir Authorize e testar cartão com o e-mail e senha do usuário.",
        "Atualizar o cartão em PUT /api/v1/metodos-pagamento/{id}.",
        "Trocar Authorize para admin / admin123.",
        "Cadastrar filme e série em POST /api/v1/conteudos.",
        "Listar catálogo em GET /api/v1/conteudos.",
        "Detalhar um conteúdo em GET /api/v1/conteudos/{id}.",
        "Adicionar favorito em POST /api/v1/favoritos.",
        "Listar favoritos em GET .../filmes e GET .../series.",
        "Opcionalmente abrir o H2 e mostrar que os dados persistiram.",
    ]:
        add_number(doc, text, decimal_num)

    add_heading(doc, "34. JSON para cadastro de usuário", 2)
    add_para(
        doc,
        "Se Maria já estiver cadastrada no seu banco, use outro e-mail e outro CPF válido. O exemplo abaixo pode "
        "ser usado em um banco limpo.",
    )
    add_code(
        doc,
        """{
  "nomeCompleto": "Ana Souza",
  "dataNascimento": "1998-08-15",
  "email": "ana@example.com",
  "senha": "Senha@123",
  "confirmarSenha": "Senha@123",
  "cpfCnpj": "11144477735",
  "numeroCartao": "4111111111111111",
  "validadeCartao": "12/2030",
  "codigoSegurancaCartao": "123",
  "nomeTitularCartao": "Ana Souza"
}""",
    )

    add_heading(doc, "35. JSON para login", 2)
    add_code(
        doc,
        """{
  "email": "maria@example.com",
  "senha": "Senha@123"
}""",
    )
    add_para(
        doc,
        "A Maria já existe no banco persistente verificado durante a revisão. Em um banco novo, faça login com "
        "o usuário que você acabou de cadastrar.",
    )

    add_heading(doc, "36. Atualização de cartão", 2)
    add_para(
        doc,
        "No Swagger, clique em Authorize e informe o e-mail e a senha do dono do cartão. Depois abra "
        "PUT /api/v1/metodos-pagamento/{id}. Use o ID retornado ou conferido no GET de métodos do usuário.",
    )
    add_code(
        doc,
        """{
  "numeroCartao": "5555555555554444",
  "validadeCartao": "12/2031",
  "codigoSegurancaCartao": "321",
  "nomeTitularCartao": "Maria da Silva"
}""",
    )

    add_heading(doc, "37. Cadastro de conteúdo como administrador", 2)
    add_para(
        doc,
        "No Swagger, abra Authorize e substitua as credenciais pelo administrador local. Depois execute "
        "POST /api/v1/conteudos.",
    )
    add_code(
        doc,
        """Username: admin
Password: admin123""",
    )
    add_code(
        doc,
        """{
  "titulo": "Coringa",
  "tipo": "FILME",
  "ano": 2019,
  "duracaoMinutos": 122,
  "relevancia": 8.5,
  "sinopse": "A origem de um dos vilões mais conhecidos dos quadrinhos.",
  "trailerUrl": "https://youtu.be/zAGVQLHvwOY",
  "genero": "Drama"
}""",
    )
    add_code(
        doc,
        """{
  "titulo": "Dark",
  "tipo": "SERIE",
  "ano": 2017,
  "duracaoMinutos": 60,
  "relevancia": 9.3,
  "sinopse": "Mistérios atravessam gerações em uma pequena cidade alemã.",
  "trailerUrl": "https://example.com/dark",
  "genero": "Ficção Científica"
}""",
    )

    add_heading(doc, "38. Favoritos e conferência", 2)
    add_para(
        doc,
        "Use os IDs realmente retornados pelo seu banco. Em um banco persistente, IDs podem não começar em 1 "
        "porque tentativas anteriores continuam registradas.",
    )
    add_code(
        doc,
        """POST /api/v1/favoritos
{
  "usuarioId": 1,
  "conteudoId": 41
}

GET /api/v1/conteudos
GET /api/v1/favoritos/usuario/1/filmes
GET /api/v1/favoritos/usuario/1/series""",
    )
    add_para(
        doc,
        "Se a resposta de séries for [], a API está funcionando: significa apenas que aquele usuário ainda não "
        "favoritou uma série.",
    )

    add_heading(doc, "39. Conferir dados pelo H2", 2)
    add_para(doc, "Abra http://localhost:8080/h2 e conecte usando a configuração do perfil dev.")
    add_code(
        doc,
        """SELECT * FROM USUARIOS;
SELECT * FROM METODO_PAGAMENTO;
SELECT * FROM CONTEUDO ORDER BY TITULO;
SELECT * FROM FAVORITOS;
SELECT * FROM PLANO;
SELECT * FROM ASSINATURA;""",
    )
    add_callout(
        doc,
        "O que observar no cartão",
        "A tabela METODO_PAGAMENTO não deve possuir número completo do cartão nem CVV. Confira ULTIMOS4, BANDEIRA, "
        "MES_EXP, ANO_EXP, NOME_PORTADOR e TOKEN_GATEWAY.",
    )

    add_heading(doc, "40. Estado do banco verificado em 01/06/2026", 2)
    add_para(
        doc,
        "Durante a revisão final, o H2 persistente continha registros válidos para demonstração e também alguns "
        "dados antigos genéricos. Eles não impedem o funcionamento, mas vale limpar antes de apresentar.",
    )
    add_table(
        doc,
        ["Situação observada", "Ação recomendada"],
        [
            ("Coringa, Dark e Breaking Bad cadastrados", "Manter para demonstração."),
            ("Interestelar duplicado", "Apagar uma cópia se quiser deixar a lista limpa."),
            ("Dois conteúdos antigos titulo=string e tipo=STRING", "Apagar antes da apresentação."),
            ("Usuário 1 possui filme favorito, mas nenhuma série", "Favoritar Dark para demonstrar as duas listas."),
        ],
        [4050, 5310],
    )
    add_callout(
        doc,
        "Importante",
        "A API agora recusa novos tipos STRING com HTTP 400. Os registros genéricos já existentes são anteriores "
        "à validação e só permanecem porque o H2 é persistente.",
        fill="FFF4D6",
        title_color=GOLD,
    )

    # Part VI
    add_heading(doc, "Parte VI - Preparação para apresentação", 1)
    add_heading(doc, "41. Roteiro oral pronto", 2)
    add_para(doc, "Você pode adaptar este texto à sua forma de falar:")
    add_callout(
        doc,
        "Abertura sugerida",
        "Meu projeto é uma API REST chamada IESPFLIX, feita com Java 17 e Spring Boot. Ela simula o backend de "
        "uma plataforma de streaming. O código foi organizado em camadas: Controller recebe as requisições, "
        "Service aplica regras de negócio, Repository acessa o banco H2, Model representa as tabelas e DTO controla "
        "os dados de entrada e saída.",
        fill=LIGHT_BLUE,
    )
    for text in [
        "No RF1, o cadastro recebe dados pessoais, confirmação de senha e cartão. Eu salvo o usuário e tokenizo o cartão na mesma transação.",
        "No RF3, a senha não vai para o banco em texto puro: eu salvo um hash BCrypt em senhaHash.",
        "No RF4, o endpoint de login busca o usuário por e-mail e usa BCrypt matches.",
        "No RF7, a rota de cartão exige HTTP Basic e confere se o cartão pertence ao usuário autenticado.",
        "No RF8, a entidade Conteudo representa filmes e séries com título, gênero, ano, duração, relevância, sinopse e trailer.",
        "No RF10, FavoritoRepository faz JOIN com Conteudo e lista FILME e SERIE separadamente.",
        "No RF11, SecurityConfig exige ROLE_ADMIN para criar, alterar ou excluir conteúdos.",
        "Também uso Bean Validation, handler global de erros, Swagger, H2 persistente, JPQL, Git, README e integração ViaCEP com OpenFeign.",
    ]:
        add_bullet(doc, text, bullet_num)

    add_heading(doc, "42. Perguntas prováveis e respostas", 2)
    add_table(
        doc,
        ["Pergunta", "Resposta curta"],
        [
            ("Por que usar DTO?", "Para controlar o contrato da API, validar entradas e evitar expor campos sensíveis."),
            ("Por que não salvar a senha?", "Salvo hash BCrypt. No login, uso matches para comparar."),
            ("Login gera token JWT?", "Não. O endpoint valida credenciais; rotas protegidas usam HTTP Basic."),
            ("Onde está RF7?", "PUT /api/v1/metodos-pagamento/{id}, com authentication.getName e conferência do dono."),
            ("Por que não salvar CVV?", "É dado sensível e desnecessário após validação. O projeto nunca persiste CVV."),
            ("Como filmes e séries são separados?", "Conteudo.tipo recebe FILME ou SERIE e FavoritoRepository filtra esse campo."),
            ("O que é JPQL?", "Consulta entidades e atributos Java; Hibernate converte para SQL."),
            ("O que é Repository?", "Interface Spring Data que abstrai acesso ao banco."),
            ("O que faz @Valid?", "Ativa Bean Validation antes de executar o service."),
            ("Onde está o serviço externo?", "ViaCepClient com @FeignClient, chamado por FuncionarioService."),
            ("Por que H2 continua com dados?", "O perfil dev usa jdbc:h2:file e ddl-auto=update."),
            ("Qual a diferença entre 401 e 403?", "401: falta autenticação válida. 403: autenticado sem permissão."),
            ("Por que existe Filme e Conteudo?", "Filme é módulo didático antigo. Conteudo é o catálogo principal do IESPFLIX."),
        ],
        [3000, 6360],
        font_size=8.5,
    )

    add_heading(doc, "43. Checklist antes de apresentar", 2)
    for text in [
        "Executar .\\mvnw.cmd test e confirmar 55 testes sem falhas.",
        "Iniciar a aplicação e abrir Swagger UI.",
        "Limpar conteúdos antigos titulo=string e duplicidade de Interestelar no H2.",
        "Cadastrar ou manter pelo menos dois filmes e duas séries.",
        "Garantir que o usuário de demonstração possui ao menos um filme e uma série favoritos.",
        "Testar POST /api/v1/auth/login.",
        "Testar Authorize com admin / admin123.",
        "Cadastrar conteúdo como admin e confirmar 201 Created.",
        "Tentar cadastrar conteúdo sem admin e explicar 401 ou 403.",
        "Abrir H2 e mostrar que senha é hash e cartão não contém número completo nem CVV.",
        "Mostrar README.md.",
        "Fazer commit e push finais no GitHub.",
    ]:
        add_bullet(doc, "☐ " + text, bullet_num)

    add_heading(doc, "44. Glossário rápido", 2)
    add_table(
        doc,
        ["Termo", "Definição"],
        [
            ("API REST", "Interface HTTP organizada por recursos, verbos e status codes."),
            ("Endpoint", "URL e verbo que executam uma operação."),
            ("JSON", "Formato textual de entrada e saída da API."),
            ("Spring Boot", "Framework que configura e executa a aplicação Java."),
            ("Bean", "Objeto criado e gerenciado pelo Spring."),
            ("Injeção de dependência", "Spring entrega objetos necessários ao construtor da classe."),
            ("Entity", "Classe Java persistida como tabela."),
            ("DTO", "Objeto usado para transportar dados pela API."),
            ("Repository", "Camada de acesso ao banco."),
            ("Service", "Camada com regras de negócio."),
            ("JPA", "Especificação Java para persistência."),
            ("Hibernate", "Implementação ORM usada para traduzir objetos e SQL."),
            ("JPQL", "Linguagem de consulta de entidades Java."),
            ("H2", "Banco leve usado localmente."),
            ("BCrypt", "Algoritmo adequado para hash de senhas."),
            ("HTTP Basic", "Autenticação que envia credenciais em cada requisição."),
            ("Swagger / OpenAPI", "Documentação interativa da API."),
            ("OpenFeign", "Cliente HTTP declarativo para APIs externas."),
            ("JaCoCo", "Ferramenta de cobertura de testes."),
        ],
        [2400, 6960],
        font_size=8.8,
    )

    add_heading(doc, "45. Resumo de bolso", 2)
    add_callout(
        doc,
        "Cinco ideias para memorizar",
        "1. Controller recebe HTTP. 2. DTO valida e protege dados. 3. Service aplica regra de negócio. "
        "4. Repository conversa com H2. 5. SecurityConfig separa rotas públicas, autenticadas e administrativas.",
        fill=LIGHT_BLUE,
    )
    add_para(
        doc,
        "Você não precisa decorar cada linha. O objetivo é conseguir explicar o caminho da informação, mostrar "
        "onde cada requisito está implementado e demonstrar o comportamento no Swagger com calma.",
        italic=True,
        color=DARK_BLUE,
        after=12,
    )
    add_para(doc, "Fim do guia.", bold=True, color=INK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, before=16, after=0)

    # Document properties
    props = doc.core_properties
    props.title = "IESPFLIX - Guia de Estudo Atualizado"
    props.subject = "Spring Boot, requisitos obrigatórios e roteiro de apresentação"
    props.author = "Samantha Hellen"
    props.keywords = "IESPFLIX, Spring Boot, API REST, Swagger, H2, estudo"
    props.comments = "Guia gerado a partir do código atual validado em 01/06/2026."

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
